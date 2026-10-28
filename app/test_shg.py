"""Self-check for the SHG creator (no browser).

Mirrors the Create SHG page's real behaviour, read off the live form:
  Save disabled -> duplicate name  |  no confirm -> never submitted
  confirm Yes   -> /shgs/list on success, stays put + toast on failure
Plus the 'Load' row filter and the fuzzy "did you mean" suggestions.

Run:  python test_shg.py
"""
import sys
sys.path.insert(0, "libs")

import os
import tempfile

import openpyxl

import config
import shg as shgmod
from shg import (ShgForm, NAME_EXISTS, SHG_LEVEL_XPATHS,
                 known_names, load_rows_only, read_shg_sheet,
                 build_shg_cascades, near_miss)

CREATE = "https://pmis/shg/create?type=Shg"
LIST = "https://pmis/shgs/list"


class FakeDriver:
    """Just enough of a WebDriver for save()."""

    def __init__(self, outcome, enabled=True, confirm="Are You Sure?",
                 redirect=True):
        self.outcome = outcome          # "saved", or the error the server toasts
        self.enabled = enabled
        self.confirm = confirm          # None = the form never asks
        self.redirect = redirect        # False = toast fires, URL lags
        self.current_url = CREATE
        self.errors = []
        self.success_toast = False
        self.confirmed = False
        self.clicked = False

    def find_elements(self, by, xpath):
        return [self] if xpath == config.SHG_FORM["save"] else []

    def is_enabled(self):
        return self.enabled

    def click(self):
        self.clicked = True

    def execute_script(self, js, *a):
        # save_state() asks for everything at once, like the real page
        return {"url": self.current_url,
                "success": self.success_toast,
                "errors": list(self.errors)}

    # --- dialog, as driver.switch_to.alert ---
    @property
    def switch_to(self):
        return self

    @property
    def alert(self):
        if self.confirm is None:
            raise Exception("no alert present")
        return self

    @property
    def text(self):
        return self.confirm

    def accept(self):
        """Answering Yes posts the form.

        On success the page toasts "Success" and then redirects; on failure it
        stays put and toasts the server's message.
        """
        self.confirmed = True
        self.confirm = None
        if self.outcome == "saved":
            self.success_toast = True
            if self.redirect:
                self.current_url = LIST
        else:
            self.errors = [self.outcome]


def save_with(**kw):
    d = FakeDriver(**kw)
    return ShgForm(d, log=lambda m: None).save(), d


def demo():
    C = config.SHG_COLUMNS
    config.SHG_SAVE_SETTLE = 0          # no need to nap in a unit check
    config.SHG_CONFIRM_TIMEOUT = 0.3
    config.SHG_SAVE_TIMEOUT = 2

    # 1. happy path: confirm answered Yes, redirected to the SHG list
    (ok, msg), d = save_with(outcome="saved")
    assert (ok, msg) == (True, "saved"), msg
    assert d.confirmed, "the Are You Sure? dialog must be accepted"

    # 2. server rejects it: stays on the form, the toast carries the reason
    (ok, msg), d = save_with(outcome="SHG Name already exists")
    assert ok is False and msg == "SHG Name already exists", (ok, msg)

    # 3. Save disabled (duplicate name caught before submit)
    (ok, msg), d = save_with(outcome="saved", enabled=False)
    assert ok is False and "already exists" in msg, (ok, msg)
    assert not d.clicked, "a disabled Save must not be clicked"

    # 4. no confirm dialog -> the form rejected itself, nothing was posted
    (ok, msg), d = save_with(outcome="saved", confirm=None)
    assert ok is False and "not submitted" in msg, (ok, msg)
    assert d.current_url == CREATE

    # 4b. the success toast reads "Success" - that is a save, not an error
    (ok, msg), d = save_with(outcome="saved", redirect=False)
    assert (ok, msg) == (True, "saved"), (ok, msg)

    # 5. known_names(): which rows must never be attempted again
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = config.SHG_SHEET
    columns = [C[k] for k in C]
    ws.append(columns + [config.SHG_STATUS_COL])
    pad = [""] * (len(columns) - 1)

    def row(name, status):
        ws.append([name] + pad + [status])

    row("ALREADY MADE", config.SHG_DONE_MARKER + " (hamlet skipped)")
    row("IS A DUPLICATE", NAME_EXISTS + ".")
    row("STILL PENDING", "")
    row("FAILED OTHERWISE", "Not filled: Geo Village")
    row("", NAME_EXISTS)                    # blank name must not join the set
    path = os.path.join(tempfile.mkdtemp(), "shg.xlsx")
    wb.save(path)

    found = known_names(read_shg_sheet(path))
    assert found == {"already made", "is a duplicate"}, found

    # 6. only rows marked Load are worked on
    rows = [
        {C["name"]: "A", config.SHG_STATUS_COL: "Load"},
        {C["name"]: "B", config.SHG_STATUS_COL: " load "},     # sloppy typing
        {C["name"]: "C", config.SHG_STATUS_COL: ""},
        {C["name"]: "D", config.SHG_STATUS_COL: "Geo OK"},
        {C["name"]: "E", config.SHG_STATUS_COL: config.SHG_DONE_MARKER},
    ]
    assert [r[C["name"]] for r in load_rows_only(rows)] == ["A", "B"]
    assert load_rows_only([{C["name"]: "A"}]) == []

    # 7. cascades come only from those rows, deduplicated
    geo = [
        {C["state"]: "ODISHA", C["district"]: "KEONJHAR", C["block"]: "GHATGAON",
         C["gp"]: "G", C["village"]: "V", C["hamlet"]: "H",
         config.SHG_STATUS_COL: "Load"},
        {C["state"]: "ODISHA", C["district"]: "KEONJHAR", C["block"]: "GHATGAON",
         C["gp"]: "G", C["village"]: "V", C["hamlet"]: "H",
         config.SHG_STATUS_COL: "Load"},                       # same cascade
        {C["state"]: "ODISHA", C["district"]: "KEONJHAR", C["block"]: "",
         C["gp"]: "G", C["village"]: "V", C["hamlet"]: "",
         config.SHG_STATUS_COL: "Load"},                       # no Block
    ]
    assert len(build_shg_cascades(load_rows_only(geo))) == 1

    # 8. a missing location suggests the near spellings, or says No Location
    class Fuzzy(ShgForm):
        OPTIONS = ["Select", "GHATGAON", "ANANDAPUR", "BANSPAL", "TELKOI"]

        def __init__(self):
            pass                       # no driver needed for suggest()

        def options(self, xpath):
            return list(self.OPTIONS)

    bot = Fuzzy()
    assert near_miss(bot, "Block", "GHATAGAON") ==         "Block not found: GHATAGAON -> GHATGAON"
    assert near_miss(bot, "Block", "ZZQQ NOWHERE") == config.NO_LOCATION
    assert "Select" not in near_miss(bot, "Block", "Selec")

    # 9. the two passes: every name is settled before anything is created
    def run(names, exists, endpoint_up=True):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = config.SHG_SHEET
        headers = [C[k] for k in ("name", "type", "partner", "workorder",
                                  "date", "state", "district", "block",
                                  "gp", "village", "hamlet", "commodity")]
        ws.append(headers + [config.SHG_STATUS_COL])
        for n in names:
            ws.append([n, config.SHG_TYPE_ALLOWED] + [""] * 10 +
                      [config.LOAD_MARKER])
        path = os.path.join(tempfile.mkdtemp(), "shg.xlsx")
        wb.save(path)

        created = []

        class Stub(ShgForm):
            def __init__(self, driver, log=print, stop_event=None):
                self.log, self.stop_event = log, stop_event
                self.opened = 0

            def _check_stop(self):
                pass

            def ensure_logged_in(self, *a):
                pass

            def verify_partner(self, *a):
                return True, "stub"

            def open_create_form(self, *a, **k):
                self.opened += 1

            def shg_name_taken(self, name):
                if not endpoint_up:
                    return None
                return str(name).strip().casefold() in exists

            def fill_shg(self, row):
                self.filling = row[C["name"]]
                return [], []

            def save(self):
                created.append(self.filling)
                return True, "saved"

        old_form, old_settings = shgmod.ShgForm, shgmod.settings
        shgmod.ShgForm = Stub
        shgmod.settings = type("S", (), {"get": staticmethod(lambda k: True)})
        try:
            n = shgmod.run_shg_creation(None, path, "u", "n", "p",
                                        log=lambda m: None)
        finally:
            shgmod.ShgForm, shgmod.settings = old_form, old_settings

        wb2 = openpyxl.load_workbook(path)
        status = [r[-1].value
                  for r in wb2[config.SHG_SHEET].iter_rows(min_row=2)]
        return n, created, status

    n, created, status = run(["FREE ONE", "TAKEN ONE", "FREE ONE"],
                             exists={"taken one"})
    # row 1 is created; row 2 is a known duplicate; row 3 repeats row 1's
    # name, so it must not be created a second time
    assert status == [config.SHG_DONE_MARKER, NAME_EXISTS, NAME_EXISTS], status
    assert created == ["FREE ONE"], created
    assert n == 1, n

    # endpoint down -> no name pass, but the run still creates
    n, created, status = run(["FREE ONE"], exists=set(), endpoint_up=False)
    assert n == 1 and status == [config.SHG_DONE_MARKER], status

    # 10. the cascade settles even when PMIS keeps resetting levels late
    config.SHG_CASCADE_SETTLE = 0.3
    config.SHG_CASCADE_TIMEOUT = 5

    class Drifty(ShgForm):
        """Resets a level the first N times it is looked at, like a late
        /districts/select reply landing after the level was chosen."""

        def __init__(self, resets, forever=None):
            self.log = lambda m: None
            self.resets = list(resets)          # levels to wipe, in order
            self.forever = forever              # a level that never sticks
            self.state = {"State": "ODISHA", "District": "KEONJHAR",
                          "Block": "GHATGAON"}
            self._cascade_set = dict(self.state)
            self.picks = []

        def current_selection(self, xpath):
            level = next(l for l, x in SHG_LEVEL_XPATHS.items() if x == xpath)
            if self.resets and self.resets[0] == level:
                self.resets.pop(0)
                self.state[level] = "Select"    # PMIS just rebuilt the list
            elif level == self.forever:
                self.state[level] = "Select"
            return self.state[level]

        def pick(self, xpath, value, timeout=None):
            level = next(l for l, x in SHG_LEVEL_XPATHS.items() if x == xpath)
            self.state[level] = value
            self.picks.append(level)
            return True

        def consume_alert(self, timeout=0.0):
            return ""

    bot = Drifty(["District", "Block", "District"])
    fixed = bot.verify_cascade()
    assert fixed == ["District", "Block", "District"], fixed
    assert bot.drifted_levels() == [], bot.drifted_levels()

    assert Drifty([]).verify_cascade() == []     # nothing to do, settles

    # a level that never sticks gives up instead of looping forever
    stubborn = Drifty([], forever="District")
    stubborn.verify_cascade()
    assert stubborn.drifted_levels(), "should still report the bad level"

    # 11. Hamlet is mandatory in PMIS and has no "none of these" entry
    class Hamlets(ShgForm):
        def __init__(self, offered):
            self.log = lambda m: None
            self.offered = offered
            self.picked = {}
            self._cascade_set = {}

        def _check_stop(self):
            pass

        def hamlet_options(self):
            return list(self.offered)

        def consume_alert(self, timeout=0.0):
            return ""

        def take_last_alert(self):
            return ""

        def current_selection(self, xpath):
            return ""

        def _report(self, *a):
            pass

        def pick(self, xpath, value, timeout=None):
            level = next(l for l, x in SHG_LEVEL_XPATHS.items() if x == xpath)
            self.picked[level] = value
            return True

        def verify_cascade(self):
            return []

    REAL = ["Select", "MUNDA SAHI", "NAYAK SAHI"]
    NONE_OPT = ["Select", "NO HAMLET", "MUNDA SAHI"]

    def cascade(offered, hamlet):
        bot = Hamlets(offered)
        row = {C["state"]: "ODISHA", C["district"]: "KEONJHAR",
               C["block"]: "GHATGAON", C["gp"]: "G", C["village"]: "V",
               C["hamlet"]: hamlet}
        failed, notes = bot.select_cascade(row, C)
        return bot.picked.get("Hamlet"), failed, notes

    # a hamlet the sheet names and PMIS has
    assert cascade(REAL, "NAYAK SAHI")[0] == "NAYAK SAHI"
    # "NO HAMLET" in the sheet -> the real "NO HAMLET" option wins
    assert cascade(NONE_OPT, "NO HAMLET")[0] == "NO HAMLET"
    # ...and where no such option exists, the first real one
    for blank in ("NO HAMLET", "NA", "", "-"):
        got, failed, notes = cascade(REAL, blank)
        assert got == "MUNDA SAHI", (blank, got)
        assert not failed and any("used 'MUNDA SAHI'" in n for n in notes), notes
    # a hamlet PMIS does not have: fall back, but say so
    got, failed, notes = cascade(REAL, "GHOST SAHI")
    assert got == "MUNDA SAHI" and not failed, (got, failed)
    assert any("GHOST SAHI" in n and "not in PMIS" in n for n in notes), notes
    # a level PMIS does not have names itself in the notes, with suggestions
    class Missing(Hamlets):
        def pick(self, xpath, value, timeout=None):
            level = next(l for l, x in SHG_LEVEL_XPATHS.items() if x == xpath)
            if level == "Village":
                return False              # PMIS has no such village
            self.picked[level] = value
            return True

        def options(self, xpath):
            return ["Select", "GARGADA BAHAL", "ROHI"]

    bot = Missing(REAL)
    failed, notes = bot.select_cascade(
        {C["state"]: "ODISHA", C["district"]: "KEONJHAR",
         C["block"]: "GHATGAON", C["gp"]: "G",
         C["village"]: "GARGADABAHAL", C["hamlet"]: ""}, C)
    assert failed == ["Village"], failed
    assert any("GARGADABAHAL" in n and "GARGADA BAHAL" in n
               for n in notes), notes

    # a village with no hamlets at all: leave the field alone
    got, failed, notes = cascade(["Select"], "ANY")
    assert got is None and not failed, (got, failed)
    assert "village has no hamlets" in notes, notes

    print("shg: all checks passed")


if __name__ == "__main__":
    demo()

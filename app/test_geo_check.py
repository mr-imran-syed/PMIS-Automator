"""Self-check for the Participant - Geo Checker (no browser).

It must behave like the SHG geo checker: run only on rows marked 'Load',
suggest near spellings for a name PMIS does not have, and write a verdict
for every row - into its OWN column, never PMIS-Update.

Run:  python test_geo_check.py
"""
import os
import sys
import tempfile

sys.path.insert(0, "libs")

import openpyxl

import config
import geo_check
from geo_check import (GeoChecker, LEVEL_ORDER, LEVEL_XPATHS, near_miss,
                       load_rows_only, validate_geo_rows)

FPO = config.PARTICIPANT_FPO_COLUMN
GEO = config.PARTICIPANT_GEO_COLUMNS
DONE = config.PARTICIPANT_DONE_MARKER
GEO_COL = config.PARTICIPANT_GEO_STATUS_COL
SHEET = "Participant Master Data"

CASCADE = {"State": "ODISHA", "District": "KEONJHAR", "Block": "GHATGAON",
           "GP": "BADAPICHULLA", "Village": "JHARBEDA"}


def workbook(rows):
    """rows: (fpo, village, pmis_update, geo_check) tuples."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    headers = [FPO] + [GEO[lvl] for lvl in LEVEL_ORDER] + ["PMIS-Update",
                                                           GEO_COL]
    ws.append(headers)
    for fpo, village, pmis, geo in rows:
        cascade = dict(CASCADE, Village=village)
        ws.append([fpo] + [cascade[lvl] for lvl in LEVEL_ORDER] +
                  [pmis, geo])
    path = os.path.join(tempfile.mkdtemp(), "participants.xlsx")
    wb.save(path)
    return path


def columns(path):
    ws = openpyxl.load_workbook(path)[SHEET]
    head = [c.value for c in ws[1]]
    out = []
    for r in ws.iter_rows(min_row=2):
        cells = dict(zip(head, [c.value for c in r]))
        out.append((cells["PMIS-Update"], cells[GEO_COL]))
    return out


class FakeChecker(GeoChecker):
    """Only the browser is faked."""

    def __init__(self, missing_villages=()):
        self.log = lambda m: None
        self.stop_event = None
        self.missing = {v.casefold() for v in missing_villages}
        self.pages = []

    def _check_stop(self):
        pass

    def goto_fpo_member_page(self, base_url, coop_id):
        self.pages.append(coop_id)
        return True

    def _check_one(self, row, retry):
        if row["Village"].casefold() in self.missing:
            return "Village"
        return None

    def options(self, xpath):          # what the dropdown offers
        return ["Select", "JHARBEDA", "JHARBEDATOWN", "BHAGAMUNDA"]


def demo():
    # 1. only rows marked Load are loaded, whatever PMIS-Update says
    rows = [
        ("697", "JHARBEDA", "", "Load"),
        ("697", "JHARBEDA", DONE, "Load"),      # already entered: still check
        ("697", "JHARBEDA", "", ""),            # not queued
        ("698", "JHARBEDA", "", config.GEO_OK_MARKER),   # checked before
    ]
    path = workbook(rows)
    headers, all_rows = geo_check._read_sheet(path, SHEET)
    assert len(load_rows_only(all_rows)) == 2, load_rows_only(all_rows)

    bot = FakeChecker()
    loaded, cascades = bot.build_location_df(path, SHEET)
    assert len(loaded) == 2 and len(cascades) == 1, (loaded, cascades)

    # 2. a good cascade -> Geo OK on its rows, PMIS-Update untouched
    _, verdicts = bot.check_locations(cascades, "http://x")
    geo_check.write_verdicts(path, SHEET, loaded, verdicts)
    assert columns(path) == [
        (None, config.GEO_OK_MARKER),
        (DONE, config.GEO_OK_MARKER),
        (None, None),                       # not queued, left alone
        (None, config.GEO_OK_MARKER),
    ], columns(path)

    # 3. a misspelled village -> the near spelling, in Geo-Check only
    path = workbook([("697", "JHARABEDA", "", "Load")])
    bot = FakeChecker(missing_villages={"JHARABEDA"})
    loaded, cascades = bot.build_location_df(path, SHEET)
    _, verdicts = bot.check_locations(cascades, "http://x")
    geo_check.write_verdicts(path, SHEET, loaded, verdicts)
    pmis, geo = columns(path)[0]
    assert pmis is None, pmis
    assert geo == "Village not found: JHARABEDA -> JHARBEDA, JHARBEDATOWN", geo

    # 4. nothing close -> No Location
    class Nowhere(FakeChecker):
        def options(self, xpath):
            return ["Select", "BHAGAMUNDA", "KANDARA"]

    path = workbook([("697", "ZZQQ NOWHERE", "", "Load")])
    bot = Nowhere(missing_villages={"ZZQQ NOWHERE"})
    loaded, cascades = bot.build_location_df(path, SHEET)
    _, verdicts = bot.check_locations(cascades, "http://x")
    geo_check.write_verdicts(path, SHEET, loaded, verdicts)
    assert columns(path)[0][1] == config.NO_LOCATION, columns(path)

    # 5. the validator refuses a run with nothing queued, and passes one
    #    whose Load rows all have a usable FPO code
    assert validate_geo_rows(workbook([("697", "JHARBEDA", "", "")]),
                             SHEET)[0] is False
    assert validate_geo_rows(workbook([("697", "JHARBEDA", "", "Load")]),
                             SHEET)[0] is True
    # a bad FPO on a row that is NOT queued must not block the run
    ok, _ = validate_geo_rows(workbook([("oops", "JHARBEDA", "", ""),
                                        ("697", "JHARBEDA", "", "Load")]),
                              SHEET)
    assert ok is True

    print("participant geo checker: all checks passed")


if __name__ == "__main__":
    demo()
